"""
pipeline/stage5_validate.py — Stage 5: 验证集预测 + Excel 位图

用模型预测 val 集, 对比真值, 生成出手位图 Excel。
后处理时在右侧追加 #N预测 列。
"""

from __future__ import annotations

import json
import logging
import os
import re
from pathlib import Path

import numpy as np
import pandas as pd
from openpyxl import load_workbook
from openpyxl.styles import PatternFill, Font

from .mq import StageRunner, StageResult

logger = logging.getLogger(__name__)

# 条件格式颜色
FILL_GREEN = PatternFill(start_color="C6EFCE", end_color="C6EFCE", fill_type="solid")
FILL_RED = PatternFill(start_color="FFC7CE", end_color="FFC7CE", fill_type="solid")
FILL_WHITE = PatternFill(start_color="FFFFFF", end_color="FFFFFF", fill_type="solid")


class Stage5Validate(StageRunner):
    name = "验证预测"
    stage_num = 5

    def run(self, config: dict, prev_data: dict) -> StageResult:
        trail_name = config["trailName"]
        dst_dir = prev_data.get("dst_dir", "")
        output_dir = prev_data.get("output_dir", "")
        mode = config.get("mode", "train")

        # ---- 加载数据 ----
        val_npy = os.path.join(dst_dir, "features_val.npy")
        labels_npy = os.path.join(dst_dir, "labels_val.npy")
        timestamps_npy = os.path.join(dst_dir, "timestamps.npy")

        if not all(os.path.exists(p) for p in [val_npy, labels_npy]):
            return StageResult(ok=False, stage=5,
                               error=f"验证数据缺失: val={os.path.exists(val_npy)}, labels={os.path.exists(labels_npy)}")

        X_val = np.load(val_npy, mmap_mode='r')
        y_val = np.load(labels_npy, mmap_mode='r')
        timestamps = np.load(timestamps_npy, allow_pickle=True)
        # 取验证集部分的时间戳
        export_meta_path = os.path.join(dst_dir, "export_meta.json")
        if os.path.exists(export_meta_path):
            with open(export_meta_path) as f:
                meta = json.load(f)
            cutoff = meta.get("cutoff", "")
            if cutoff:
                ts_str = np.array([str(t) for t in timestamps])
                val_start = np.searchsorted(ts_str, cutoff)
                timestamps = timestamps[val_start:]

        if len(timestamps) > len(X_val):
            timestamps = timestamps[:len(X_val)]
        elif len(timestamps) < len(X_val):
            X_val = X_val[:len(timestamps)]
            y_val = y_val[:len(timestamps)]

        logger.info(f"验证数据: X={X_val.shape}, y={y_val.shape}, ts={len(timestamps)}")

        # ---- 加载模型 ----
        model_path = prev_data.get("model_path", "")
        if not model_path or not os.path.exists(model_path):
            return StageResult(ok=False, stage=5, error=f"模型文件不存在: {model_path}")

        import joblib
        model = joblib.load(model_path)

        # ---- 预测 ----
        y_prob = model.predict_proba(X_val)[:, 1]
        y_pred = np.full_like(y_prob, -1, dtype=int)
        y_pred[y_prob >= 0.7] = 1
        y_pred[y_prob <= 0.3] = 0

        # 对比真值
        y_true = y_val.copy()
        y_true[y_true < 0] = -1

        # 计算成功率
        has_label = y_true >= 0
        has_signal = y_pred >= 0
        both = has_label & has_signal

        n_total = len(y_prob)
        n_signals = int(has_signal.sum())
        n_correct = int((both & (y_pred == y_true)).sum())
        signal_rate = n_signals / n_total if n_total > 0 else 0
        correct_rate = n_correct / n_signals if n_signals > 0 else 0

        logger.info(f"预测统计: 信号率={signal_rate:.1%} ({n_signals}/{n_total}), "
                     f"正确率={correct_rate:.1%} ({n_correct}/{n_signals})")

        # ---- 生成/更新 Excel ----
        os.makedirs(output_dir, exist_ok=True)
        excel_path = os.path.join(output_dir, "03_live_predictions.xlsx")

        # 构建当前预测数据
        df_new = pd.DataFrame({
            "timestamp": [str(t) for t in timestamps[:len(y_prob)]],
            "p_up": np.round(y_prob, 4),
            "prediction": [
                "涨" if p == 1 else ("跌" if p == 0 else "震荡")
                for p in y_pred
            ],
        })

        if mode == "train" or not os.path.exists(excel_path):
            # 首次创建
            df_full = pd.DataFrame({
                "timestamp": [str(t) for t in timestamps[:len(y_prob)]],
                "close": [0.0] * len(y_prob),
                "actual_direction": [
                    "涨" if y == 1 else ("跌" if y == 0 else "NA")
                    for y in y_true[:len(y_prob)]
                ],
                "#1预测": [
                    f"{p:.3f}{'涨' if p>=0.7 else '跌' if p<=0.3 else '震'}"
                    for p in y_prob
                ],
            })
            df_full.to_excel(excel_path, index=False, engine="openpyxl")
            logger.info(f"Excel 已创建: {excel_path}")
        else:
            # 后处理: 追加列
            self._append_column(excel_path, df_new, y_prob)
            logger.info(f"Excel 已追加列: {excel_path}")

        return StageResult(ok=True, stage=5,
                           message=f"信号率={signal_rate:.1%} 正确率={correct_rate:.1%}",
                           data={"n_total": n_total, "n_signals": n_signals,
                                 "n_correct": n_correct,
                                 "signal_rate": round(signal_rate, 4),
                                 "correct_rate": round(correct_rate, 4),
                                 "excel_path": os.path.abspath(excel_path)})

    def _append_column(self, excel_path: str, df_new: pd.DataFrame, y_prob: np.ndarray):
        """向后处理 Excel 追加预测列"""
        wb = load_workbook(excel_path)
        ws = wb.active

        # 计算下一列编号
        col_num = ws.max_column + 1
        col_header = f"#{col_num - 3}预测"  # 减去 timestamp/close/actual_direction 三列

        # 写表头
        ws.cell(row=1, column=col_num, value=col_header)
        ws.cell(row=1, column=col_num).font = Font(bold=True)

        # 写数据
        for i, prob in enumerate(y_prob):
            row = i + 2  # 第1行是表头
            val = f"{prob:.3f}{'涨' if prob>=0.7 else '跌' if prob<=0.3 else '震'}"
            cell = ws.cell(row=row, column=col_num, value=val)

            # 条件格式
            if prob >= 0.7:
                cell.fill = FILL_GREEN
            elif prob <= 0.3:
                cell.fill = FILL_RED

        wb.save(excel_path)
        wb.close()

    def _count_existing_pred_columns(self, excel_path: str) -> int:
        """计算已有预测列数"""
        if not os.path.exists(excel_path):
            return 0
        wb = load_workbook(excel_path)
        ws = wb.active
        count = 0
        for col in range(1, ws.max_column + 1):
            header = str(ws.cell(row=1, column=col).value or "")
            if re.match(r'^#\d+预测$', header):
                count += 1
        wb.close()
        return count
